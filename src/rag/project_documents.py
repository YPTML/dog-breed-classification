"""Project-wide readers. Document contents are data, never executable instructions."""
from __future__ import annotations

import csv
import fnmatch
import hashlib
import json
import os
import re
import zipfile
import stat
from collections import Counter
from pathlib import Path
from typing import Iterator

from src.rag.knowledge_base import DocumentChunk, compute_text_hash

TEXT = {'.md', '.markdown', '.txt', '.rst', '.py', '.js', '.ts', '.tsx', '.jsx',
        '.html', '.htm', '.css', '.yaml', '.yml', '.toml', '.ini', '.cfg', '.tex',
        '.bib', '.bat', '.ps1', '.sh', '.sql', '.log', '.xml', '.m'}
STRUCTURED = {'.json', '.jsonl', '.csv', '.tsv', '.ipynb'}
DOCUMENTS = {'.pdf', '.docx'}
ARTIFACTS = {'.keras', '.npy', '.npz'}
IMAGE = {'.jpg', '.jpeg', '.png', '.webp', '.gif', '.bmp', '.tif', '.tiff'}
PRUNED = {'.git', '.venv', 'venv', 'env', 'node_modules', '__pycache__', '.cache',
          '.pytest_cache', '.mypy_cache', '.ruff_cache', '.ipynb_checkpoints',
          '_audit_tmp', 'tmp', 'temp', 'cache', 'htmlcov'}
SECRET = re.compile(r'(^\.env($|\.)|credentials|service[-_]account|private[-_]key|secrets?\.)', re.I)
SECRET_VALUE = re.compile(
    r'(?im)((?:api[_-]?key|access[_-]?token|refresh[_-]?token|password|client_secret|authorization)'
    r'[\s\"\x27]*[:=][\s\"\x27]*)([^\s\"\x27,;\}\]]{8,})')


def redact(text: str) -> str:
    text = re.sub(r'-----BEGIN [^-]*PRIVATE KEY-----.*?-----END [^-]*PRIVATE KEY-----',
                  '[REDACTED PRIVATE KEY]', text, flags=re.S)
    text = re.sub(r'(?i)([?&](?:key|api_key|token|access_token)=)[^&\s]+',r'\1[REDACTED]',text)
    text = re.sub(r'(?i)(Bearer\s+)[A-Za-z0-9._~+/=-]{12,}',r'\1[REDACTED]',text)
    return SECRET_VALUE.sub(r'\1[REDACTED]', text)


def permitted(config, relative: str, *, metadata: bool = False, trusted_walk: bool = False) -> bool:
    # Reject irrelevant binary formats before doing filesystem resolution.
    if Path(relative).suffix.lower() not in TEXT | STRUCTURED | DOCUMENTS | ARTIFACTS:
        return False
    root = config.project_root.absolute()
    path = root / relative
    if Path(relative).is_absolute() or '..' in Path(relative).parts:
        return False
    try:
        resolved = path.absolute() if trusted_walk else path.resolve()
        resolved.relative_to(root)
    except (OSError, ValueError):
        return False
    if (not trusted_walk and path.is_symlink()) or any(part in PRUNED for part in Path(relative).parts):
        return False
    index_dir = config.knowledge_base.index_dir.absolute()
    if resolved == index_dir or index_dir in resolved.parents:
        return False
    if SECRET.search(path.name) or path.suffix.lower() in {'.key', '.pem', '.p12', '.pfx'}:
        return False
    normalized = relative.replace('\\', '/')
    # Exclude the project's generated index/reports even when the active DB is outside it.
    if normalized.startswith('outputs/rag_index/'):
        return False
    for pattern in config.knowledge_base.exclude_patterns:
        # Binary artifacts have metadata readers; their byte contents stay excluded.
        if metadata and path.suffix.lower() in ARTIFACTS and pattern in {'*.keras', '*.npy', '*.npz'}:
            continue
        if fnmatch.fnmatchcase(normalized, pattern) or fnmatch.fnmatchcase(path.name, pattern):
            return False
        if pattern.endswith('/**') and normalized.startswith(pattern[:-3] + '/'):
            return False
    if config.knowledge_base.discover_project:
        return path.suffix.lower() in TEXT | STRUCTURED | DOCUMENTS | ARTIFACTS
    return normalized in {p.replace('\\', '/') for p in config.knowledge_base.allowlist_files}


def discover(config) -> tuple[list[str], dict]:
    """List readable sources and disclose non-text coverage without opening image bytes."""
    root = config.project_root.absolute()
    index_dir = config.knowledge_base.index_dir.absolute()
    sources = []
    inventory = {'scanned_files': 0, 'by_extension': {}, 'image_folders': {},
                 'excluded_directories': [], 'unsupported_files': [], 'excluded_files': []}
    extensions = Counter()
    if not config.knowledge_base.discover_project:
        candidates = list(config.knowledge_base.allowlist_files)
    else:
        candidates = []
        for base, dirs, files in os.walk(root, followlinks=False):
            keep = []
            for name in sorted(dirs):
                p = Path(base) / name
                # Windows junctions can point to other projects or cycle.
                is_link = p.is_symlink() or bool(getattr(p.lstat(), 'st_file_attributes', 0) & getattr(stat,'FILE_ATTRIBUTE_REPARSE_POINT',0))
                if name in PRUNED or is_link or p == index_dir:
                    inventory['excluded_directories'].append(p.relative_to(root).as_posix())
                else:
                    keep.append(name)
            dirs[:] = keep
            for name in sorted(files):
                rel = (Path(base) / name).relative_to(root).as_posix()
                ext = Path(name).suffix.lower()
                inventory['scanned_files'] += 1
                extensions[ext or '(no extension)'] += 1
                if ext in IMAGE:
                    folder = Path(rel).parent.as_posix()
                    inventory['image_folders'][folder] = inventory['image_folders'].get(folder, 0) + 1
                elif ext in TEXT | STRUCTURED | DOCUMENTS | ARTIFACTS and permitted(config, rel, metadata=True, trusted_walk=True):
                    candidates.append(rel)
                elif ext in TEXT | STRUCTURED | DOCUMENTS | ARTIFACTS:
                    inventory['excluded_files'].append(rel)
                else:
                    inventory['unsupported_files'].append(rel)
    for rel in candidates:
        if permitted(config, rel, metadata=True, trusted_walk=config.knowledge_base.discover_project) and (root / rel).is_file():
            sources.append(rel)
    inventory['by_extension'] = dict(extensions)
    inventory['candidate_files'] = len(sources)
    inventory['image_note'] = 'Ảnh được thống kê theo thư mục và tra cứu qua manifest/nhãn. Chưa phân tích nội dung thị giác.'
    return sorted(set(sources)), inventory


class ProjectReader:
    def __init__(self, config):
        self.config = config
        self.max_chars = max(400, config.chunking.markdown_max_tokens * 5)
        self.issues: list[dict] = []

    def chunk(self, rel, content, title, locator, doc_type='text', start=0, end=0, **metadata):
        content = redact(content)
        return DocumentChunk(
            chunk_id=f'{rel}#{locator}', source_path=rel, title=title,
            content=content, start_line=start, end_line=end, doc_type=doc_type,
            content_hash=compute_text_hash(content), version=self.config.knowledge_base.version,
            metadata={'locator': locator, **metadata})

    def split_text(self, rel, text, title=None, doc_type='text', locator_prefix='', **metadata):
        """Bound every chunk, including long single lines; never discard the tail."""
        title = title or rel
        lines = text.splitlines()
        buffer = []
        start = 1
        length = 0
        section = title
        serial = 0
        for lineno, line in enumerate(lines, 1):
            heading = re.match(r'^#{1,6}\s+(.+)', line)
            if buffer and (length + len(line) > self.max_chars or heading):
                serial += 1
                yield self.chunk(rel, '\n'.join(buffer), section,
                                 f'{locator_prefix}L{start}-L{lineno-1}-c{serial}',
                                 doc_type, start, lineno-1, **metadata)
                buffer, length, start = [], 0, lineno
            if heading:
                section = f'{title} > {heading.group(1)}'
            if not buffer:
                start = lineno
            if len(line) > self.max_chars:
                for offset in range(0, len(line), self.max_chars):
                    serial += 1
                    yield self.chunk(rel, line[offset:offset+self.max_chars], section,
                                     f'{locator_prefix}L{lineno}-c{serial}', doc_type,
                                     lineno, lineno, **metadata)
            else:
                buffer.append(line)
                length += len(line) + 1
        if buffer and any(x.strip() for x in buffer):
            serial += 1
            yield self.chunk(rel, '\n'.join(buffer), section,
                             f'{locator_prefix}L{start}-L{len(lines)}-c{serial}',
                             doc_type, start, len(lines), **metadata)

    def json_leaves(self, data, pointer=''):
        if isinstance(data, dict) and data:
            for key, value in data.items():
                token = str(key).replace('~', '~0').replace('/', '~1')
                yield from self.json_leaves(value, pointer + '/' + token)
        elif isinstance(data, list) and data:
            for i, value in enumerate(data):
                yield from self.json_leaves(value, pointer + '/' + str(i))
        else:
            yield pointer or '/', data

    def read(self, rel: str) -> Iterator[DocumentChunk]:
        if not permitted(self.config, rel, metadata=True):
            return
        path = self.config.project_root / rel
        suffix = path.suffix.lower()
        if suffix not in ARTIFACTS and path.stat().st_size > self.config.knowledge_base.max_file_bytes:
            self.issues.append({'source_path': rel, 'status': 'too_large', 'bytes': path.stat().st_size})
            return
        if suffix in TEXT:
            text = path.read_text(encoding='utf-8-sig', errors='replace')
            if '\ufffd' in text:
                self.issues.append({'source_path': rel, 'status': 'encoding_replacements'})
            kind = 'code' if suffix in {'.py','.js','.ts','.tsx','.jsx'} else 'text'
            if 'papers' in path.parts and path.name.endswith('.notes.md'):
                kind = 'paper_note'
            yield from self.split_text(rel, text, doc_type=kind)
        elif suffix in {'.csv', '.tsv'}:
            with path.open(encoding='utf-8-sig', newline='') as stream:
                reader = csv.reader(stream, delimiter='\t' if suffix == '.tsv' else ',')
                header = next(reader, [])
                count = 0
                for count, row in enumerate(reader, 1):
                    content = '\n'.join(f'{key}: {value}' for key, value in zip(header, row))
                    if len(row) > len(header):
                        content += '\nextra_columns: ' + json.dumps(row[len(header):], ensure_ascii=False)
                    yield from self.split_text(rel, content, title=f'{rel} | row {count}',
                                               doc_type='structured_csv', locator_prefix=f'row={count}&',
                                               row=count, columns=header)
                yield self.chunk(rel, f'Table: {rel}\nRows: {count}\nColumns: {", ".join(header)}',
                                 f'{rel} > Table schema', 'schema', 'structured_csv', row_count=count, columns=header)
        elif suffix in {'.json', '.jsonl'}:
            with path.open(encoding='utf-8-sig') as stream:
                values = ((f'line={i}&', json.loads(line)) for i,line in enumerate(stream,1) if line.strip()) if suffix == '.jsonl' else [('',json.load(stream))]
                for prefix, data in values:
                    # Group small leaves while retaining exact JSON pointers and all values.
                    buffer, length, first, last = [], 0, '', ''
                    for pointer, value in self.json_leaves(data):
                        entry = f'{pointer}: {json.dumps(value, ensure_ascii=False)}'
                        if buffer and length + len(entry) > self.max_chars:
                            yield self.chunk(rel, '\n'.join(buffer), f'{rel} > {first}',
                                             prefix+'json='+first, 'structured_json', json_pointer=first, last_pointer=last)
                            buffer, length = [], 0
                        if not buffer: first = pointer
                        last = pointer
                        if len(entry) > self.max_chars:
                            yield from self.split_text(rel, entry, f'{rel} > {pointer}', 'structured_json',
                                                       prefix+'json='+pointer+'&', json_pointer=pointer)
                        else:
                            buffer.append(entry); length += len(entry)+1
                    if buffer:
                        yield self.chunk(rel, '\n'.join(buffer), f'{rel} > {first}', prefix+'json='+first,
                                         'structured_json', json_pointer=first, last_pointer=last)
        elif suffix == '.ipynb':
            data = json.loads(path.read_text(encoding='utf-8-sig'))
            for i, cell in enumerate(data.get('cells', []), 1):
                text = cell.get('source', [])
                text = ''.join(text) if isinstance(text,list) else str(text)
                for output in cell.get('outputs', []):
                    out = output.get('text', output.get('data', {}).get('text/plain', ''))
                    text += '\n' + (''.join(out) if isinstance(out,list) else str(out))
                yield from self.split_text(rel,text,f'{rel} > cell {i}','notebook',f'cell={i}&',cell=i)
        elif suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted and not reader.decrypt(''):
                raise ValueError('PDF encrypted: password required')
            meta = reader.metadata or {}
            title = str(meta.get('/Title') or path.stem)
            author = str(meta.get('/Author') or '')
            sidecar = path.with_suffix('.metadata.json')
            if sidecar.is_file():
                try:
                    if not permitted(self.config,sidecar.relative_to(self.config.project_root).as_posix()) or sidecar.stat().st_size > 1024*1024:
                        raise ValueError('Metadata sidecar is excluded or too large')
                    annotated = json.loads(sidecar.read_text(encoding='utf-8-sig'))
                    if not isinstance(annotated,dict):
                        raise ValueError('Metadata must be an object')
                    title = str(annotated.get('title') or title)
                    authors = annotated.get('authors',[])
                    if isinstance(authors,list) and all(isinstance(a,str) for a in authors):
                        author = ', '.join(authors) or author
                except (OSError,ValueError):
                    self.issues.append({'source_path':rel,'status':'invalid_pdf_metadata','metadata_file':sidecar.name})
            for page_no, page in enumerate(reader.pages, 1):
                text = page.extract_text(extraction_mode='layout') or ''
                if len(text.strip()) < 20:
                    self.issues.append({'source_path': rel, 'page': page_no, 'status': 'ocr_required'})
                    continue
                yield from self.split_text(rel,text,f'{title} > page {page_no}','paper',f'page={page_no}&',
                                           page=page_no, document_title=title, author=author,
                                           extraction='pdf_text', visual_content_analyzed=False)
        elif suffix == '.docx':
            from docx import Document
            doc = Document(path)
            parts = [p.text for p in doc.paragraphs]
            parts += [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
            yield from self.split_text(rel,'\n'.join(parts),doc_type='document',extraction='docx_text')
        elif suffix == '.keras':
            with zipfile.ZipFile(path) as archive:
                for name in ['metadata.json','config.json']:
                    if name in archive.namelist():
                        info = archive.getinfo(name)
                        if info.file_size > 10*1024*1024:
                            self.issues.append({'source_path': rel, 'status':'artifact_metadata_too_large','member':name})
                            continue
                        text = archive.read(name).decode('utf-8')
                        yield from self.split_text(rel,text,f'{rel} > {name}','artifact_metadata',name+'&',
                                                   metadata_only=True, archive_member=name)
        elif suffix in {'.npy','.npz'}:
            import numpy as np
            if suffix == '.npy':
                array = np.load(path, mmap_mode='r', allow_pickle=False)
                text = f'Array: {rel}\nshape: {array.shape}\ndtype: {array.dtype}\nOnly header metadata; values not indexed.'
                yield self.chunk(rel,text,rel,'metadata','artifact_metadata',metadata_only=True)
                del array
            else:
                with zipfile.ZipFile(path) as archive:
                    names = archive.namelist()
                yield from self.split_text(rel,'\n'.join(names),rel,'artifact_metadata','members&',metadata_only=True)
